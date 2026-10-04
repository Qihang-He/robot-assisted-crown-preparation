#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Supplementary Tables Generation

Generate a single Excel file (Supplementary_Tables.xlsx) containing all
supplementary tables as separate worksheets. Properly formatted with
bold headers, auto‑adjusted column widths, and descriptive titles.

Copyright © 2026 Qihang He
Released under the MIT licence; see LICENSE.

Paper step: supplementary tables (Table S1 and the per-specimen tables behind the main tables).
Inputs:  analysis_results/ (all CSV files produced by the analysis scripts)
Outputs: Supplementary_Tables.xlsx (one worksheet per table)
Licence: MIT (see LICENSE).
"""

import os, glob, warnings
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter
from datetime import datetime

warnings.filterwarnings('ignore')

# ============================ Configuration ============================
ANALYSIS_DIR = r"path/to/your/project/data/analysis_results"
OUTPUT_FILE  = r"path/to/your/output/Supplementary_Tables.xlsx"

# =================== TABLE DEFINITIONS =====================
TABLE_DEFS = [
    ("PSR per specimen",
     "preparation_accuracy/PSR_summary.csv",
     "Table S1",
     "Per‑specimen PSR error (μ_base, σ_base) estimated from the non‑operative base region (F0). "
     "Supporting data for the main text (Appendix Table 4)."),

    ("Preparation accuracy per specimen",
     "preparation_accuracy/region_results.csv",
     "Table S2",
     "Full per‑specimen, per‑region preparation accuracy after PSR correction. "
     "Includes μ_prep, σ_prep, R_pos, μ_pos, P75. Supporting data for the main text."),

    ("Regional summary – preparation accuracy",
     "preparation_accuracy/region_summary.csv",
     "Table S3",
     "Regional summary of preparation accuracy after PSR correction. Same as main text Table 2."),

    ("Between‑specimen repeatability",
     "preparation_accuracy/consistency_PSR.csv",
     "Table S4",
     "Between‑specimen SD of μ_prep for each region. Supports Section 3.1.3."),

    ("Crown accuracy per specimen",
     "crown_accuracy/crown_accuracy.csv",
     "Table S5",
     "Per‑specimen crown printing accuracy (μ_crown, σ_crown). Supporting data for the main text."),

    ("Crown accuracy – group summary",
     "crown_accuracy/crown_group_summary.csv",
     "Table S6",
     "Group‑wise summary of crown printing accuracy. Supporting data for the main text."),

    ("Internal fit per specimen",
     "internal_fit/gap_results.csv",
     "Table S7",
     "Per‑specimen, per‑region internal fit (μ_gap, σ_gap). Supporting data for the main text."),

    ("Internal fit – group summary",
     "internal_fit/gap_group_summary.csv",
     "Table S8",
     "Group‑wise summary of internal fit. Supporting data for the main text."),

    ("Marginal gap per specimen",
     "marginal_gap/marginal_gap_results.csv",
     "Table S9",
     "Per‑specimen marginal gap metrics (μ_marginal, σ_marginal, P75, IoU_XY). "
     "Supporting data for the main text."),

    ("Marginal gap – group summary",
     "marginal_gap/group_summary.csv",
     "Table S10",
     "Group‑wise summary of marginal gap. Supporting data for the main text."),
]

# Pairwise p‑value tables
PVAL_PATTERNS = [
    ("marginal_gap_visualization", "pairwise_pvalues_*.csv"),
    ("internal_fit_visualization", "pairwise_pvalues_*.csv"),
    ("crown_accuracy_visualization", "pairwise_pvalues_*.csv"),
]


def set_column_width(ws):
    """Auto‑fit column widths"""
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            try:
                val = str(cell.value) if cell.value is not None else ""
                max_len = max(max_len, len(val))
            except:
                pass
        ws.column_dimensions[col_letter].width = min(max_len + 3, 60)


def insert_header(ws, title, description, ncol):
    """Insert two header rows at the top with title and description."""
    ws.insert_rows(1, 2)
    title_cell = ws.cell(row=1, column=1, value=title)
    title_cell.font = Font(bold=True, size=14)
    desc_cell = ws.cell(row=2, column=1, value=description)
    desc_cell.font = Font(size=10, italic=True)
    if ncol > 1:
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncol)
    for cell in ws[3]:
        cell.font = Font(bold=True)


def main():
    print("=" * 60)
    print("  Generating Supplementary_Tables.xlsx")
    print("  Output:", OUTPUT_FILE)
    print("=" * 60)

    if not os.path.exists(ANALYSIS_DIR):
        print(f"ERROR: analysis_results directory not found at {ANALYSIS_DIR}")
        return

    writer = pd.ExcelWriter(OUTPUT_FILE, engine='openpyxl')

    for sheet_name, src_path, label, desc in TABLE_DEFS:
        full_path = os.path.join(ANALYSIS_DIR, src_path)
        if not os.path.exists(full_path):
            print(f"  [WARN] {full_path} not found, skipping.")
            continue
        df = pd.read_csv(full_path)
        short_name = label.replace(" ", "_")[:31]
        df.to_excel(writer, sheet_name=short_name, index=False)
        ws = writer.sheets[short_name]
        insert_header(ws, f"{label}: {sheet_name}", desc, df.shape[1])
        set_column_width(ws)
        print(f"  {label} -> sheet '{short_name}' ({len(df)} rows)")

    # Pairwise p‑values
    pval_list = []
    for folder, pattern in PVAL_PATTERNS:
        full_pattern = os.path.join(ANALYSIS_DIR, folder, pattern)
        files = glob.glob(full_pattern)
        for fpath in files:
            try:
                p_df = pd.read_csv(fpath)
                metric = os.path.basename(fpath).replace('pairwise_pvalues_', '').replace('.csv', '')
                p_df.insert(0, 'Metric', metric)
                pval_list.append(p_df)
            except Exception as e:
                print(f"  [ERROR] reading {fpath}: {e}")
    if pval_list:
        pval_all = pd.concat(pval_list, ignore_index=True)
        sheet_name = "Pairwise_p_values"
        pval_all.to_excel(writer, sheet_name=sheet_name, index=False)
        ws = writer.sheets[sheet_name]
        insert_header(ws, "Table S11: Pairwise comparisons (Bonferroni‑corrected p‑values)",
                      "All pairwise Mann‑Whitney U test p‑values after Bonferroni correction.",
                      pval_all.shape[1])
        set_column_width(ws)
        print(f"  Table S11 -> sheet '{sheet_name}' ({len(pval_all)} rows)")

    writer.close()
    print(f"\nTotal sheets: {len(pd.ExcelFile(OUTPUT_FILE).sheet_names)}")
    print("Done.")


if __name__ == '__main__':
    main()